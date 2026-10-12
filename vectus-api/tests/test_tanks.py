"""Tank lookup against the Excel itself.

An independent oracle re-reads ``data/Tank_details.xlsx`` with its own, simpler
parser. Every exhaustive test asks the lookup a question and checks the answer
against what the oracle says the sheet contains — so a bug in the lookup's
parsing or filtering cannot also hide in the expected values.
"""

from __future__ import annotations

import json
import os
import re
from itertools import product
from pathlib import Path

import openpyxl
import pytest

from app import tanks as T

SHEET = Path(__file__).resolve().parent.parent / "data" / "Tank_details.xlsx"
BRANDS = ["Vectus", "Ganga", "Waterwell"]
SPECIAL = {"chemical tank", "isi tank", "underground", "group"}


# --- the oracle ---------------------------------------------------------------------


def _oracle_rows():
    ws = openpyxl.load_workbook(SHEET, data_only=True).worksheets[0]
    rows = [r for r in ws.iter_rows(values_only=True) if r[1]]
    out = []
    for r in rows[1:]:
        name = r[1].strip()
        first, rest = name.split(" ", 1)
        if first.lower() == "water" and rest.lower().startswith("well "):
            first, rest = "Waterwell", rest[5:]
        brand = {"vectus": "Vectus", "ganga": "Ganga", "waterwell": "Waterwell"}[first.lower()]
        caps = sorted({int(x.replace(",", ""))
                       for x in re.findall(r"\d{1,3}(?:,\d{3})+(?!\d)|\d+", str(r[4]))})
        out.append({
            "brand": brand,
            "key": re.sub(r"[^a-z0-9]+", " ", rest.lower()).strip(),
            "layers": int(r[2]),
            "colours": {c.strip().title() for c in str(r[3]).split(",") if c.strip()},
            "caps": caps,
            "guarantee": int(r[6]) if isinstance(r[6], (int, float)) else None,
            "mould": str(r[5]).strip().title(),
        })
    return out


ROWS = _oracle_rows()
ALL_CAPS = sorted({c for r in ROWS for c in r["caps"]})
NOT_MADE = [50, 120, 250, 400, 600, 800, 900, 1200, 1250, 1800, 2500, 4000, 7000,
            12000, 40000, 100000]
COLOURS = sorted({c for r in ROWS for c in r["colours"]})
LAYERS = sorted({r["layers"] for r in ROWS})


def oracle(brand, **f):
    """Names the sheet says match, with the home/special rule the bot follows."""
    pool = [r for r in ROWS if brand in (None, r["brand"])]

    def ok(r):
        if "cap" in f and f["cap"] not in r["caps"]:
            return False
        if "colour" in f and f["colour"] not in r["colours"]:
            return False
        if "layers" in f and r["layers"] != f["layers"]:
            return False
        if "guarantee" in f and r["guarantee"] != f["guarantee"]:
            return False
        return True

    hits = [r for r in pool if ok(r)]
    regular = [r for r in hits if r["key"] not in SPECIAL]
    return {f"{r['brand']} {r['key']}" for r in no_chemical(regular or hits)}


def no_chemical(rows):
    """The Chemical Tank is named only when nothing else matches (run 849)."""
    rest = [r for r in rows if r["key"] != "chemical tank"]
    return rest or rows


@pytest.fixture(scope="module")
def cat():
    return T.load_catalog(SHEET)


def names(result) -> set[str]:
    """All tanks a found answer covers (shown + 'more'), normalised like the oracle."""
    shown = [t["name"] for t in result.get("tanks", [])] + result.get("more", [])
    return {f"{n.split(' ', 1)[0]} {re.sub(r'[^a-z0-9]+', ' ', n.split(' ', 1)[1].lower()).strip()}"
            for n in shown}


def card_matches_sheet(card) -> bool:
    brand, model = card["name"].split(" ", 1)
    key = re.sub(r"[^a-z0-9]+", " ", model.lower()).strip()
    row = next(r for r in ROWS if r["brand"] == brand and r["key"] == key)
    return (card["capacities"] == row["caps"] and set(card["colours"]) == row["colours"]
            and card["layers"] == row["layers"] and card["guarantee_years"] == row["guarantee"]
            and card["mould"] == row["mould"])


def every_card(result):
    cards = list(result.get("tanks", []))
    cards += result.get("alternatives", [])
    cards += result.get("other_brands", [])
    for key in ("tank", "same_tank"):
        if key in result:
            cards.append(result[key])
    return cards


def assert_truthful(result):
    """No answer may carry a number the sheet does not have for that tank."""
    for card in every_card(result):
        assert card_matches_sheet(card), card
    assert len(result.get("tanks", [])) <= T.MAX_RESULTS
    json.dumps(result)  # always serialisable


# --- the sheet loads as the oracle reads it -----------------------------------------


def test_every_row_loads_exactly(cat):
    assert len(cat.tanks) == len(ROWS) == 39
    for t, r in zip(cat.tanks, ROWS):
        assert (t.brand, t.key) == (r["brand"], r["key"])
        assert list(t.capacities) == r["caps"]
        assert set(t.colours) == r["colours"]
        assert (t.layers, t.guarantee, t.mould) == (r["layers"], r["guarantee"], r["mould"])


@pytest.mark.parametrize("name,caps", [
    ("Vectus Cool", [550, 750, 1100, 1500, 2000, 3000, 5000]),
    ("Vectus Safe", [550, 750, 1100, 1500, 2000, 3000, 5000]),
    ("Vectus Granito", [550, 750, 1100, 1500, 2000]),
    ("Vectus Silk", [500, 750, 1000, 1500, 2000]),
    ("Vectus Loft", [100, 150, 200, 300, 500, 1000]),
    ("Vectus TenX", [500, 750, 1000, 1500, 2000, 3000, 5000]),
    ("Vectus Chemical Tank", [550, 750, 1100, 1500, 2000, 3000, 5000, 10000]),
    ("Vectus Group", [15000, 20000, 25000, 30000]),     # "15,000 , 20,000": thousands
    ("Waterwell T-90", [500, 750, 1000, 1500, 2000]),    # "Water Well T-90"
    ("Waterwell TenX", [500, 750, 1000, 1500, 2000, 3000, 5000]),  # "WaterWell TenX"
    ("Vectus Mini", [1000]),
])
def test_values_where_the_old_prompt_was_wrong(cat, name, caps):
    t = next(t for t in cat.tanks if t.name == name)
    assert list(t.capacities) == caps


def test_features_leave_out_the_sheets_self_contradictions(cat):
    horizontal = next(t for t in cat.tanks if t.name == "Vectus Horizontal Tank")
    assert not any("three-layer" in f.lower() for f in horizontal.features)
    for t in cat.tanks:
        assert not any(re.search(r"\d+ to [\d,]+ litres", f) for f in t.features), t.name


def test_layer_names_are_read(cat):
    cool = next(t for t in cat.tanks if t.name == "Vectus Cool")
    assert len(cool.layer_names) == 5 and "Microban" in cool.layer_names[-1]
    t90 = next(t for t in cat.tanks if t.name == "Vectus T-90")  # "Layer4- Extra" in the sheet
    assert len(t90.layer_names) == 6


# --- exhaustive: every brand × every capacity ---------------------------------------


@pytest.mark.parametrize("brand,cap", list(product(BRANDS + [None], ALL_CAPS + NOT_MADE)))
def test_capacity_only(cat, brand, cap):
    r = T.find(cat, brand=brand or "all", capacity=cap)
    assert_truthful(r)
    expected = oracle(brand, cap=cap)
    if expected:
        assert r["status"] == "found" and names(r) == expected
    else:
        assert r["status"] == "no_match"
        pool = [x for x in ROWS if brand in (None, x["brand"])]
        regular = [x for x in pool if x["key"] not in SPECIAL]
        caps = sorted({c for x in regular for c in x["caps"]})
        if not any(c > cap for c in caps):  # bigger than every home tank: large ones count
            caps = sorted({c for x in pool for c in x["caps"]})
        lower = max([c for c in caps if c < cap], default=None)
        higher = min([c for c in caps if c > cap], default=None)
        assert r["closest_capacities"] == {"lower": lower, "higher": higher}
        for option in r["options"]:  # every tank offered really has that capacity
            for n in option["tanks"]:
                b, m = n.split(" ", 1)
                row = next(x for x in ROWS if x["brand"] == b
                           and x["key"] == re.sub(r"[^a-z0-9]+", " ", m.lower()).strip())
                assert option["capacity"] in row["caps"]


@pytest.mark.parametrize("brand,colour", list(product(BRANDS + [None], COLOURS)))
def test_colour_only(cat, brand, colour):
    r = T.find(cat, brand=brand or "all", colour=colour)
    assert_truthful(r)
    expected = oracle(brand, colour=colour)
    if expected:
        assert r["status"] == "found" and names(r) == expected
    else:
        assert r["status"] == "no_match"


@pytest.mark.parametrize("brand,cap,colour",
                         list(product(BRANDS, ALL_CAPS + [1200, 4000], COLOURS)))
def test_capacity_and_colour(cat, brand, cap, colour):
    r = T.find(cat, brand=brand, capacity=cap, colour=colour)
    assert_truthful(r)
    expected = oracle(brand, cap=cap, colour=colour)
    if expected:
        assert r["status"] == "found" and names(r) == expected
    else:
        assert r["status"] == "no_match"
        for option in r.get("options", []):
            for n in option["tanks"]:  # the colour asked is kept in every suggestion
                b, m = n.split(" ", 1)
                row = next(x for x in ROWS if x["brand"] == b
                           and x["key"] == re.sub(r"[^a-z0-9]+", " ", m.lower()).strip())
                assert colour in row["colours"] and option["capacity"] in row["caps"]


@pytest.mark.parametrize("brand,layers", list(product(BRANDS, LAYERS + [7, 8, 12])))
def test_layers_only(cat, brand, layers):
    r = T.find(cat, brand=brand, layers=layers)
    assert_truthful(r)
    expected = oracle(brand, layers=layers)
    assert (r["status"] == "found" and names(r) == expected) if expected else \
        r["status"] == "no_match"


@pytest.mark.parametrize("brand,years", list(product(BRANDS, [10, 15, 5, 20])))
def test_guarantee_only(cat, brand, years):
    r = T.find(cat, brand=brand, guarantee=years)
    assert_truthful(r)
    expected = oracle(brand, guarantee=years)
    assert (r["status"] == "found" and names(r) == expected) if expected else \
        r["status"] == "no_match"


@pytest.mark.parametrize("brand,cap,layers", list(product(BRANDS, ALL_CAPS, LAYERS)))
def test_capacity_and_layers(cat, brand, cap, layers):
    r = T.find(cat, brand=brand, capacity=cap, layers=layers)
    assert_truthful(r)
    expected = oracle(brand, cap=cap, layers=layers)
    if expected:
        assert r["status"] == "found" and names(r) == expected
    else:
        assert r["status"] == "no_match"


# --- exhaustive: every model × every capacity / colour ------------------------------


MODEL_ROWS = [r for r in ROWS]


@pytest.mark.parametrize("row,cap", [(r, c) for r in MODEL_ROWS for c in ALL_CAPS + NOT_MADE],
                         ids=lambda v: v if isinstance(v, int) else f"{v['brand']}-{v['key']}")
def test_named_model_and_capacity(cat, row, cap):
    r = T.find(cat, brand=row["brand"], model=row["key"], capacity=cap)
    assert_truthful(r)
    if cap in row["caps"]:
        assert r["status"] == "found" and len(r["tanks"]) == 1
        assert r["message"].startswith(f"Yes, {row['brand']} ")
    else:
        assert r["status"] == "not_available"
        lower = max([c for c in row["caps"] if c < cap], default=None)
        higher = min([c for c in row["caps"] if c > cap], default=None)
        assert r["closest_capacities"] == {"lower": lower, "higher": higher}
        assert f"does not come in {cap} litre" in r["message"]
        for alt in r.get("alternatives", []):
            assert cap in alt["capacities"] and alt["name"].startswith(row["brand"])


@pytest.mark.parametrize("row,colour", [(r, c) for r in MODEL_ROWS for c in COLOURS],
                         ids=lambda v: v if isinstance(v, str) else f"{v['brand']}-{v['key']}")
def test_named_model_and_colour(cat, row, colour):
    r = T.find(cat, brand=row["brand"], model=row["key"], colour=colour)
    assert_truthful(r)
    if colour in row["colours"]:
        assert r["status"] == "found"
    else:
        assert r["status"] == "not_available" and "comes only in" in r["message"]
        for alt in r.get("alternatives", []):
            assert colour in alt["colours"]


@pytest.mark.parametrize("row", MODEL_ROWS, ids=lambda r: f"{r['brand']}-{r['key']}")
def test_named_model_alone_gives_its_full_details(cat, row):
    r = T.find(cat, brand=row["brand"], model=row["key"])
    assert r["status"] == "found" and len(r["tanks"]) == 1
    card = r["tanks"][0]
    assert card_matches_sheet(card) and "layer_names" in card and "features" in card


# --- a model the brand does not sell ------------------------------------------------


@pytest.mark.parametrize("brand,model,twin", [
    ("Vectus", "Pro", "Vectus Smart"), ("Vectus", "Health", "Vectus Safe"),
    ("Vectus", "Italian", "Vectus Cool"), ("Ganga", "Safe", "Ganga Health"),
    ("Ganga", "Smart", "Ganga Pro"), ("Ganga", "Italian", "Ganga Cool"),
    ("Waterwell", "Cool", "Waterwell Italian"), ("Waterwell", "Safe", "Waterwell Health"),
    ("Waterwell", "Smart", "Waterwell Pro"),
])
def test_same_tank_under_the_brands_own_name(cat, brand, model, twin):
    r = T.find(cat, brand=brand, model=model)
    assert r["status"] == "not_in_brand" and r["same_tank"]["name"] == twin
    assert_truthful(r)


@pytest.mark.parametrize("brand,model,available", [
    ("Ganga", "Mini", ["Vectus Mini"]), ("Waterwell", "Mini", ["Vectus Mini"]),
    ("Ganga", "Chemical", ["Vectus Chemical Tank"]), ("Ganga", "Group", ["Vectus Group"]),
])
def test_model_only_another_brand_sells(cat, brand, model, available):
    r = T.find(cat, brand=brand, model=model)
    assert r["status"] == "not_in_brand"
    assert [c["name"] for c in r["other_brands"]] == available


# --- brand rule ---------------------------------------------------------------------


@pytest.mark.parametrize("state", ["Tamil Nadu", "TN", "tamilnadu", "Kerala", "KL", "Karnataka",
                                   "KA", "Andhra Pradesh", "AP", "Andhra", "Telangana", "TS"])
def test_south_states_default_to_ganga(cat, state):
    r = T.find(cat, state=state, capacity=1100)
    assert r["brand"] == "Ganga" and r["brand_reason"] == "state default"
    assert all(t["name"].startswith("Ganga") for t in r["tanks"])


@pytest.mark.parametrize("state", ["Uttar Pradesh", "UP", "Delhi", "Bihar", "Punjab",
                                   "Maharashtra", "Jammu and Kashmir", "West Bengal"])
def test_other_states_default_to_vectus(cat, state):
    assert T.find(cat, state=state, capacity=1100)["brand"] == "Vectus"


def test_no_brand_no_state_is_vectus(cat):
    assert T.find(cat, capacity=1100)["brand"] == "Vectus"


def test_a_named_brand_beats_the_state(cat):
    assert T.find(cat, brand="Vectus", state="Tamil Nadu", capacity=1100)["brand"] == "Vectus"
    assert T.find(cat, brand="Waterwell", state="UP", capacity=1100)["brand"] == "Waterwell"


def test_brand_inside_the_model_text_wins(cat):
    r = T.find(cat, brand="Vectus", model="Ganga Pro")
    assert r["brand"] == "Ganga" and r["tanks"][0]["name"] == "Ganga Pro"


@pytest.mark.parametrize("spelling", ["vectus", "Victus", "vektus", "VECTUS", "vector"])
def test_brand_spellings(cat, spelling):
    assert T.find(cat, brand=spelling, capacity=1000)["brand"] == "Vectus"


def test_water_well_spelling(cat):
    assert T.find(cat, brand="water well", capacity=1000)["brand"] == "Waterwell"


def test_a_competitor_brand_is_flagged_not_used(cat):
    r = T.find(cat, brand="Sintex", capacity=1000)
    assert r["brand"] == "Vectus" and r["ignored"] == {"brand": "Sintex"}


# --- what callers say, as the bot passes it -----------------------------------------


@pytest.mark.parametrize("said,expected", [
    (1100, (1100, 1100)), (1100.0, (1100, 1100)), ("1100", (1100, 1100)),
    ("1,100", (1100, 1100)), ("1100 litre", (1100, 1100)), ("1100L", (1100, 1100)),
    ("1100 ltr", (1100, 1100)), ("1.1k", (1100, 1100)), ("2k", (2000, 2000)),
    ("gyarah sau", (1100, 1100)), ("saade saat sau", (750, 750)),
    ("saade paanch sau", (550, 550)), ("paanch sau", (500, 500)), ("saat sau", (700, 700)),
    ("ek hazaar", (1000, 1000)), ("dedh hazaar", (1500, 1500)), ("do hazaar", (2000, 2000)),
    ("dhai hazaar", (2500, 2500)), ("teen hazaar", (3000, 3000)),
    ("paanch hazaar", (5000, 5000)), ("das hazaar", (10000, 10000)),
    ("pachees hazaar", (25000, 25000)), ("one thousand", (1000, 1000)),
    ("fifteen hundred", (1500, 1500)), ("1000 se 1500", (1000, 1500)),
    ("1000-1500", (1000, 1500)), ("1000 to 2000 litre", (1000, 2000)),
    ("1 se 2 hazaar", (1000, 2000)), ("15,000", (15000, 15000)),
])
def test_capacity_as_spoken(said, expected):
    cap = T.parse_capacity(said)
    assert (cap.low, cap.high) == expected and cap.values is None


@pytest.mark.parametrize("said,values", [
    ("1000 ya 1500", (1000, 1500)), ("1000 or 1500", (1000, 1500)),
    ("1000/1500", (1000, 1500)), ("1000, 1500", (1000, 1500)),
    ("750 ya 1000 ya 1500", (750, 1000, 1500)), ("ek ya do hazaar", (1000, 2000)),
])
def test_capacity_as_a_choice(said, values):
    assert T.parse_capacity(said).values == values


@pytest.mark.parametrize("said,low,high", [
    ("500 tak", 1, 500), ("upto 500", 1, 500), ("500 se kam", 1, 500), ("below 1000", 1, 1000),
    ("2000 se upar", 2000, T.NO_LIMIT), ("2000 se zyada", 2000, T.NO_LIMIT),
    ("above 5000", 5000, T.NO_LIMIT), ("5000+", 5000, T.NO_LIMIT),
    ("kam se kam 1000", 1000, T.NO_LIMIT), ("1000 se 1500 tak", 1000, 1500),
])
def test_capacity_open_ranges(said, low, high):
    cap = T.parse_capacity(said)
    assert (cap.low, cap.high, cap.values) == (low, high, None)


def test_choice_does_not_match_sizes_in_between(cat):
    # "1000 ya 1500" must not offer a tank only because it has 1100.
    r = T.find(cat, capacity="1000 ya 1500")
    assert names(r) == oracle("Vectus", cap=1000) | oracle("Vectus", cap=1500)


def test_up_to_and_at_least(cat):
    assert names(T.find(cat, capacity="300 tak")) == {"Vectus loft"}
    big = T.find(cat, capacity="5000 se upar")
    assert names(big) == {"Vectus chemical tank", "Vectus isi tank", "Vectus group"} or         {"Vectus cool", "Vectus safe"} <= names(big)


@pytest.mark.parametrize("junk", ["", "abc", "bada wala", 0, -500, "null", True, "10"])
def test_capacity_that_is_not_a_capacity(junk):
    assert T.parse_capacity(junk) is None


@pytest.mark.parametrize("said,model", [
    ("cool", "Vectus Cool"), ("Cool", "Vectus Cool"), ("vectus cool", "Vectus Cool"),
    ("kool", "Vectus Cool"), ("vector school", "Vectus Cool"), ("vectors school", "Vectus Cool"),
    ("school tank", "Vectus Cool"), ("cooler", "Vectus Cool"), ("कूल", "Vectus Cool"),
    ("tenx", "Vectus TenX"), ("ten x", "Vectus TenX"), ("10x", "Vectus TenX"),
    ("Ten-X", "Vectus TenX"), ("टेन एक्स", "Vectus TenX"),
    ("T-90", "Vectus T-90"), ("t90", "Vectus T-90"), ("T 90", "Vectus T-90"),
    ("t ninety", "Vectus T-90"),
    ("granito", "Vectus Granito"), ("granite", "Vectus Granito"), ("grenito", "Vectus Granito"),
    ("puff", "Vectus Puff"), ("puf", "Vectus Puff"), ("पफ", "Vectus Puff"),
    ("silk", "Vectus Silk"), ("smart", "Vectus Smart"), ("safe", "Vectus Safe"),
    ("loft", "Vectus Loft"), ("mini", "Vectus Mini"), ("isi", "Vectus ISI tank"),
    ("ISI tank", "Vectus ISI tank"), ("underground", "Vectus Underground"),
    ("under ground tank", "Vectus Underground"), ("horizontal", "Vectus Horizontal Tank"),
    ("chemical", "Vectus Chemical Tank"), ("group", "Vectus Group"),
    ("Cool wala tank", "Vectus Cool"), ("puff model chahiye", "Vectus Puff"),
    ("water tank puff", "Vectus Puff"),
    ("granitto", "Vectus Granito"), ("smaart", "Vectus Smart"),  # typos
])
def test_model_names_as_heard(cat, said, model):
    r = T.find(cat, model=said)
    assert r["status"] == "found", (said, r)
    assert r["tanks"][0]["name"] == model


@pytest.mark.parametrize("said", ["sintex", "xyz", "supreme", "plastic drum", "aquatech"])
def test_unknown_model_lists_the_real_ones(cat, said):
    r = T.find(cat, model=said)
    assert r["status"] == "unknown_model" and "Cool" in r["models"]


@pytest.mark.parametrize("said,colour", [
    ("white", "White"), ("safed", "White"), ("सफेद", "White"), ("green", "Green"),
    ("hara", "Green"), ("blue", "Blue"), ("neela", "Blue"), ("नीला", "Blue"),
    ("black", "Black"), ("kala", "Black"), ("yellow", "Lime Yellow"), ("peela", "Lime Yellow"),
    ("lime yellow", "Lime Yellow"), ("ruby", "Ruby"), ("emerald", "Emerald"), ("pearl", "Pearl"),
    ("GREEN colour", "Green"), ("hara rang", "Green"),
])
def test_colours_as_heard(said, colour):
    assert T.parse_colours(said)[0] == {colour}


def test_two_colours(cat):
    r = T.find(cat, colour="green ya blue")
    assert names(r) == oracle("Vectus", colour="Green") | oracle("Vectus", colour="Blue")


@pytest.mark.parametrize("said", ["red", "laal", "grey", "orange", "pink"])
def test_unknown_colour_lists_the_real_ones(cat, said):
    r = T.find(cat, colour=said)
    assert r["status"] == "unknown_colour"
    assert set(r["colours"]) == {c for x in ROWS if x["brand"] == "Vectus" for c in x["colours"]}


@pytest.mark.parametrize("said,layers", [("5", 5), (5, 5), ("5 layer", 5), ("paanch", 5),
                                         ("10 layers", 10), ("teen layer", 3)])
def test_layers_as_heard(cat, said, layers):
    r = T.find(cat, layers=said)
    assert all(t["layers"] == layers for t in r["tanks"])


@pytest.mark.parametrize("said", ["15 saal", "15 years", "15", 15, "15 yrs"])
def test_guarantee_as_heard(cat, said):
    r = T.find(cat, guarantee=said)
    assert r["status"] == "found" and all(t["guarantee_years"] == 15 for t in r["tanks"])


def test_no_guarantee(cat):
    r = T.find(cat, guarantee="no guarantee")
    assert names(r) == {f"Vectus {x['key']}" for x in no_chemical(
        [x for x in ROWS if x["brand"] == "Vectus" and x["guarantee"] is None])}


@pytest.mark.parametrize("said,need", [
    ("garmi mein thanda", "cooling"), ("cooling", "cooling"), ("heat", "cooling"),
    ("dhoop", "cooling"), ("underground", "underground"), ("zameen ke neeche", "underground"),
    ("kam jagah", "small_space"), ("parapet", "small_space"), ("commercial", "commercial"),
    ("chemical", "chemical"), ("bacteria", "hygiene"),
])
def test_needs_as_heard(said, need):
    assert T.parse_need(said)[0] == need


def test_cooling_need_comes_from_the_layer_names(cat):
    r = T.find(cat, need="cooling")
    for n in names(r):
        tank = next(t for t in cat.tanks if f"{t.brand} {t.key}" == n)
        text = " ".join(tank.layer_names).lower()
        assert "foam" in text or "heat-reflective" in text or "weather-resistant" in text
    assert "Vectus cool" in names(r) and "Vectus puff" in names(r) and "Vectus tenx" in names(r)
    assert "Vectus smart" not in names(r)


def test_cooling_with_capacity(cat):
    r = T.find(cat, need="thanda pani", capacity=1100)
    assert names(r) == {"Vectus cool", "Vectus granito"}


def test_underground_need(cat):
    r = T.find(cat, need="underground", capacity=2000)
    assert r["status"] == "found" and r["tanks"][0]["name"] == "Vectus Underground"


def test_underground_size_not_made(cat):
    r = T.find(cat, need="underground", capacity=3000)
    assert r["status"] == "not_available"
    assert r["closest_capacities"] == {"lower": 2000, "higher": 5000}


def test_commercial_opens_the_large_tanks(cat):
    r = T.find(cat, need="commercial", capacity=20000)
    assert r["tanks"][0]["name"] == "Vectus Group"
    r = T.find(cat, need="commercial", capacity=5000)
    assert {"Vectus chemical tank", "Vectus isi tank", "Vectus cool"} <= names(r)


# --- home vs special tanks ----------------------------------------------------------


def test_home_enquiry_does_not_offer_chemical_or_isi(cat):
    r = T.find(cat, capacity=1100)
    assert names(r) == {"Vectus safe", "Vectus cool", "Vectus granito"}


def test_special_tanks_answer_when_only_they_match(cat):
    assert names(T.find(cat, capacity=10000)) == {"Vectus isi tank"}
    assert names(T.find(cat, capacity=25000)) == {"Vectus group"}
    assert "Vectus underground" in names(T.find(cat, colour="black"))
    assert names(T.find(cat, layers=1)) == {"Vectus chemical tank"}
    assert names(T.find(cat, capacity=200)) == {"Vectus loft"}  # ISI 200 L is special


# --- the whole phrase in "model" ----------------------------------------------------


@pytest.mark.parametrize("phrase,status,first", [
    ("Cool 1100 litre", "found", "Vectus Cool"),
    ("Granito 3000", "not_available", None),
    ("T-90 1000", "found", "Vectus T-90"),
    ("green safe", "found", "Vectus Safe"),
    ("blue puff", "not_available", None),
    ("5 layer wala", "found", "Vectus Cool"),
    ("10 layer", "found", None),
    ("water tank", "catalogue", None),
])
def test_phrases_in_the_model_field(cat, phrase, status, first):
    r = T.find(cat, model=phrase)
    assert r["status"] == status, r
    if first:
        assert r["tanks"][0]["name"] == first
    assert_truthful(r)


def test_explicit_capacity_beats_a_number_in_the_model(cat):
    r = T.find(cat, model="Cool 1100", capacity=750)
    assert r["status"] == "found" and "750 litre" in r["message"]


@pytest.mark.parametrize("phrase,expected", [
    ("Cool aur Puff", ["Vectus Cool", "Vectus Puff"]),
    ("Cool vs Puff", ["Vectus Cool", "Vectus Puff"]),
    ("Smart, Safe and Cool", ["Vectus Smart", "Vectus Safe", "Vectus Cool"]),
    ("Pro ya Health", ["Vectus Smart", "Vectus Safe"]),  # Vectus's own names
])
def test_comparisons(cat, phrase, expected):
    r = T.find(cat, model=phrase)
    assert r["status"] == "found" and [t["name"] for t in r["tanks"]] == expected
    assert all("layer_names" in t for t in r["tanks"])
    assert r["message"].startswith("Compare:")
    assert_truthful(r)


def test_comparison_in_ganga(cat):
    r = T.find(cat, model="Cool aur Safe", state="TN")
    assert [t["name"] for t in r["tanks"]] == ["Ganga Cool", "Ganga Health"]


# --- nothing usable -----------------------------------------------------------------


def test_nothing_given_asks_for_capacity(cat):
    r = T.find(cat)
    assert r["status"] == "catalogue" and "capacity" in r["message"]
    assert "Vectus Chemical Tank" not in r["models"]


@pytest.mark.parametrize("kwargs", [
    {"capacity": "null", "colour": "none", "model": "", "layers": "NA"},
    {"capacity": "abc"},
    {"model": "  ", "colour": ""},
])
def test_blank_or_junk_inputs(cat, kwargs):
    r = T.find(cat, **kwargs)
    assert r["status"] == "catalogue"


def test_junk_capacity_is_reported(cat):
    assert T.find(cat, capacity="bahut bada")["ignored"] == {"capacity": "bahut bada"}


def test_unknown_need_is_reported_and_ignored(cat):
    r = T.find(cat, capacity=1100, need="sasta")
    assert r["status"] == "found" and r["ignored"] == {"need": "sasta"}


def test_relaxation_names_what_fails(cat):
    r = T.find(cat, layers=10, colour="green")
    assert r["status"] == "no_match"
    assert {x["without"] for x in r["relaxed"]} == {"layers", "colour"}


def test_many_matches_are_capped_and_listed(cat):
    r = T.find(cat, capacity=1000)
    assert len(r["tanks"]) == T.MAX_RESULTS
    assert names(r) == oracle("Vectus", cap=1000)
    assert "more" in r and "ask which" in r["message"]


def test_messages_name_only_real_tanks(cat):
    real = {t.name for t in cat.tanks}
    for kwargs in ({"capacity": 1100}, {"colour": "green"}, {"model": "Smart", "capacity": 1100},
                   {"capacity": 1200}, {"need": "cooling"}, {"model": "Pro"},
                   {"layers": 7}, {"model": "Cool aur Puff"}):
        r = T.find(cat, **kwargs)
        for n in re.findall(r"\b(?:Vectus|Ganga|Waterwell) [A-Z][\w-]*(?: Tank| tank)?", r["message"]):
            assert n in real or n.rsplit(" ", 1)[0] in real, (kwargs, n)


# --- startup refuses a bad sheet ----------------------------------------------------


def _write_sheet(tmp_path, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Sr No.", "Tank Name", "No. of layers", "Color", "Capacities", "Type",
               "Gurantee (Yrs)", "Features", "Name of layers"])
    for row in rows:
        ws.append(row)
    path = tmp_path / "t.xlsx"
    wb.save(path)
    return path


GOOD = [1, "Vectus Cool", 5, "White", "550, 750", "Roto", 15, "1-A: b", "1-X 2-Y"]


def test_a_good_minimal_sheet_loads(tmp_path):
    assert len(T.load_catalog(_write_sheet(tmp_path, [GOOD])).tanks) == 1


@pytest.mark.parametrize("bad,why", [
    ([1, "Sintex Cool", 5, "White", "550", "Roto", 15, "", ""], "brand"),
    ([1, "Vectus Cool", 5, "White", "", "Roto", 15, "", ""], "capacities"),
    ([1, "Vectus Cool", "five", "White", "550", "Roto", 15, "", ""], "layers"),
    ([1, "Vectus Cool", 5, "", "550", "Roto", 15, "", ""], "colour"),
    ([1, "Vectus Cool", 5, "White", "550", "Injection", 15, "", ""], "Roto or Blow"),
])
def test_a_bad_row_stops_startup(tmp_path, bad, why):
    with pytest.raises(T.CatalogError, match=why):
        T.load_catalog(_write_sheet(tmp_path, [bad]))


def test_a_duplicate_row_stops_startup(tmp_path):
    with pytest.raises(T.CatalogError, match="twice"):
        T.load_catalog(_write_sheet(tmp_path, [GOOD, GOOD]))


def test_a_missing_column_stops_startup(tmp_path):
    wb = openpyxl.Workbook()
    wb.active.append(["Tank Name", "Color"])
    wb.active.append(["Vectus Cool", "White"])
    path = tmp_path / "t.xlsx"
    wb.save(path)
    with pytest.raises(T.CatalogError, match="column"):
        T.load_catalog(path)


def test_a_missing_file_stops_startup(tmp_path):
    with pytest.raises(T.CatalogError):
        T.load_catalog(tmp_path / "nope.xlsx")


# --- HTTP ---------------------------------------------------------------------------

os.environ.setdefault("VECTUS_API_KEY", "test-key")


@pytest.fixture(scope="module")
def client():
    from fastapi.testclient import TestClient

    from app import main
    main.API_KEY = "test-key"
    with TestClient(main.app) as c:
        yield c


H = {"X-API-Key": "test-key"}


def test_http_found(client):
    r = client.post("/tanks", json={"capacity": 1100, "colour": "green"}, headers=H)
    assert r.status_code == 200 and r.json()["tanks"][0]["name"] == "Vectus Safe"


def test_http_trailing_slash_and_color_spelling(client):
    r = client.post("/tanks/", json={"color": "peela"}, headers=H)
    assert r.status_code == 200 and names(r.json()) == {"Vectus t 90", "Vectus tenx"}


def test_http_strings_for_numbers(client):
    r = client.post("/tanks", json={"capacity": "1100 litre", "layers": "5",
                                    "guarantee": "15 saal"}, headers=H)
    assert r.json()["tanks"][0]["name"] == "Vectus Cool"


def test_http_llm_nulls(client):
    r = client.post("/tanks", json={"model": "Cool", "capacity": "null", "colour": "",
                                    "state": "None", "brand": None}, headers=H)
    assert r.json()["status"] == "found"


def test_http_unknown_fields_ignored(client):
    r = client.post("/tanks", json={"model": "Cool", "product": "Water tank", "x": 1}, headers=H)
    assert r.json()["status"] == "found"


def test_http_empty_body(client):
    assert client.post("/tanks", json={}, headers=H).json()["status"] == "catalogue"


def test_http_very_long_text(client):
    r = client.post("/tanks", json={"model": "cool " * 300}, headers=H)
    assert r.status_code == 200


@pytest.mark.parametrize("headers", [{}, {"X-API-Key": "wrong"}])
def test_http_auth(client, headers):
    assert client.post("/tanks", json={"model": "Cool"}, headers=headers).status_code == 401


def test_http_health_counts_tanks(client):
    assert client.get("/health").json()["tanks"] == 39


def test_http_reply_stays_small(client):
    """The bot reads the whole reply; the biggest one must stay compact."""
    worst = 0
    for body in ({"capacity": 1000}, {"model": "Smart, Safe, Cool and TenX"},
                 {"need": "commercial"}, {"capacity": 1200}, {"model": "Smart", "capacity": 1100}):
        worst = max(worst, len(client.post("/tanks", json=body, headers=H).text))
    assert worst < 6000


# --- the spoken message carries only the sheet's numbers ----------------------------

_DESCRIBED = re.compile(
    r"((?:Vectus|Ganga|Waterwell) [^:;.]+?): (\d+) layer, ([^,]+), "
    r"(?:(\d+) litre only|\d+ to \d+ litres \(([\d, ]+)\)), (no guarantee|\d+ year guarantee)")


def assert_message_truthful(result):
    for name, layers, colours, only, caps, guarantee in _DESCRIBED.findall(result["message"]):
        brand, model = name.split(" ", 1)
        key = re.sub(r"[^a-z0-9]+", " ", model.lower()).strip()
        row = next(r for r in ROWS if r["brand"] == brand and r["key"] == key)
        said = [int(only)] if only else [int(c) for c in caps.split(", ")]
        assert said == row["caps"], (name, said)
        assert int(layers) == row["layers"]
        assert set(colours.split(" / ")) == row["colours"]
        assert guarantee == (f"{row['guarantee']} year guarantee" if row["guarantee"]
                             else "no guarantee")


def test_random_requests_never_misstate_the_sheet(cat):
    """5000 random mixes of every input, including junk, on every brand."""
    import random

    rnd = random.Random(845)
    models = [None, "cool", "puff", "smart", "safe", "granito", "t90", "tenx", "silk",
              "loft", "mini", "pro", "health", "italian", "isi", "underground", "group",
              "chemical", "horizontal", "Cool aur Puff", "sintex", "green safe", "Cool 1100"]
    caps = [None, *ALL_CAPS, *NOT_MADE, "dedh hazaar", "1000 se 1500", "abc"]
    colours = [None, *COLOURS, "peela", "hara", "red", "green ya blue"]
    seen = set()
    for _ in range(5000):
        kwargs = {
            "brand": rnd.choice([None, *BRANDS, "all", "Sintex"]),
            "state": rnd.choice([None, "TN", "UP", "Kerala", "Delhi"]),
            "model": rnd.choice(models), "capacity": rnd.choice(caps),
            "colour": rnd.choice(colours), "layers": rnd.choice([None, *LAYERS, 7]),
            "guarantee": rnd.choice([None, 10, 15, "no", 20]),
            "need": rnd.choice([None, "cooling", "underground", "commercial", "sasta"]),
        }
        r = T.find(cat, **kwargs)
        seen.add(r["status"])
        assert_truthful(r)
        assert_message_truthful(r)
        assert r["message"]
    assert seen >= {"found", "not_available", "no_match", "not_in_brand", "unknown_model",
                    "unknown_colour"}
